"""Research summary DOCX, CSV/ZIP output; figures embedded without Kaleido."""
from __future__ import annotations
from io import BytesIO
from zipfile import ZipFile,ZIP_DEFLATED
from datetime import datetime
import re
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches,Pt
from .auto_analysis import auto_profile
from .preprocess import detect_column_roles,quality_table


def csv_bytes(df:pd.DataFrame)->bytes:
    return df.to_csv(index=False).encode('utf-8-sig')


def zip_tables(tables:dict[str,pd.DataFrame])->bytes:
    stream=BytesIO()
    with ZipFile(stream,'w',ZIP_DEFLATED) as z:
        for key,df in tables.items():
            safe=re.sub(r'[^\w-]','_',key.lower())
            z.writestr(f'{safe}.csv',csv_bytes(df))
    return stream.getvalue()


def assemble_tables(df:pd.DataFrame,roles=None,run=None,changes=None)->dict[str,pd.DataFrame]:
    roles=roles or detect_column_roles(df)
    profile=auto_profile(df,roles)
    tabs={'column_quality':quality_table(df,roles),'numeric_summary':profile['numeric'],
          'correlation_pairs':profile['pairs']}
    for col,frequency in profile['categorical'].items():tabs[f'frequency_{col}']=frequency
    if run is not None:
        tabs.update({'regression_cv':run.cv_metrics,'regression_holdout':run.final_metrics,
                     'regression_predictions':run.predictions})
    if changes:tabs['data_preparation_log']=pd.DataFrame(changes)
    return tabs


def _add_table(doc:Document,df:pd.DataFrame,max_rows:int=25)->None:
    if df.empty:
        doc.add_paragraph('No applicable results.');return
    data=df.head(max_rows)
    t=doc.add_table(rows=1,cols=len(data.columns));t.style='Light Shading Accent 1'
    for i,c in enumerate(data.columns):t.rows[0].cells[i].text=str(c)
    for row in data.itertuples(index=False,name=None):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if isinstance(v,(float,np.floating)) and np.isfinite(v):v=f'{v:.4g}'
            cells[i].text=str(v)
    if len(df)>max_rows:doc.add_paragraph(f'Only the first {max_rows} of {len(df)} rows are shown; download the table ZIP for complete results.')


def _figure_png(df:pd.DataFrame,run=None)->bytes|None:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    fig,ax=plt.subplots(figsize=(7.5,3.8))
    if run is not None:
        p=run.predictions
        ax.scatter(p['observed'],p['predicted'],alpha=.7)
        lo=min(p['observed'].min(),p['predicted'].min());hi=max(p['observed'].max(),p['predicted'].max())
        ax.plot([lo,hi],[lo,hi],'--',color='gray')
        ax.set_xlabel('Observed');ax.set_ylabel('Predicted');ax.set_title('Held-out regression predictions')
    else:
        numeric=df.select_dtypes(include=['number']).columns
        if not len(numeric):plt.close(fig);return None
        x=pd.to_numeric(df[numeric[0]],errors='coerce').dropna()
        ax.hist(x,bins=min(24,max(5,int(np.sqrt(len(x))))))
        ax.set_xlabel(str(numeric[0]));ax.set_ylabel('Frequency');ax.set_title(f'Distribution of {numeric[0]}')
    fig.tight_layout();buf=BytesIO();fig.savefig(buf,format='png',dpi=145);plt.close(fig);buf.seek(0)
    return buf.getvalue()


def make_docx(df:pd.DataFrame,filename:str,raw_shape:tuple,changes:list[dict],roles=None,run=None,lang:str='en')->bytes:
    roles=roles or detect_column_roles(df)
    profile=auto_profile(df,roles)
    def L(en,th): return th if lang=='th' else en
    doc=Document();sec=doc.sections[0];sec.header_distance=Inches(.4)
    normal=doc.styles['Normal'];normal.font.name='TH Sarabun New' if lang=='th' else 'Aptos';normal.font.size=Pt(11 if lang=='th' else 9)
    doc.add_heading('Interactive Data Lab',level=0)
    doc.add_paragraph(L('Automated tabular research summary | Lecture-aligned, no LLM','สรุปการวิเคราะห์ข้อมูลตารางอัตโนมัติ | สอดคล้องกับบทเรียน | ไม่ใช้ LLM'))
    doc.add_paragraph(L('Generated: ','สร้างเมื่อ: ')+datetime.now().strftime('%Y-%m-%d %H:%M'))
    doc.add_heading(L('1. Data provenance and representation','1. แหล่งที่มาและการแทนข้อมูล'),1)
    doc.add_paragraph(L(f'Dataset: {filename}. Original shape: {raw_shape[0]} rows × {raw_shape[1]} variables. Analysis shape: {len(df)} rows × {len(df.columns)} variables.',f'ชุดข้อมูล: {filename} · ต้นฉบับ: {raw_shape[0]} แถว × {raw_shape[1]} ตัวแปร · ชุดวิเคราะห์: {len(df)} แถว × {len(df.columns)} ตัวแปร'))
    doc.add_paragraph(L(f'Detected roles (with user overrides): {len(roles.numeric)} numeric, {len(roles.categorical)} categorical, {len(roles.datetime)} datetime, {len(roles.id_like)} identifiers.',f'ประเภทตัวแปร (รวมที่ผู้ใช้ปรับ): เชิงตัวเลข {len(roles.numeric)} หมวดหมู่ {len(roles.categorical)} เวลา {len(roles.datetime)} ตัวระบุ {len(roles.id_like)}'))
    doc.add_paragraph(L('Column roles are heuristic and should be verified against the data dictionary. Raw uploads are not overwritten.','ประเภทตัวแปรถูกตรวจเบื้องต้นโดยกฎอัตโนมัติ จึงควรตรวจสอบกับคำอธิบายข้อมูล ไม่เขียนทับไฟล์ต้นฉบับ'))
    doc.add_heading(L('2. Preparation and audit','2. การเตรียมข้อมูลและประวัติการแก้ไข'),1)
    if changes:
        for item in changes:doc.add_paragraph(f'{item.get("Action","")}: {item.get("Details","")}',style='List Bullet')
    else:doc.add_paragraph(L('No user-approved data transformations were applied.','ไม่มีการแก้ไขข้อมูลที่ได้รับการยืนยันจากผู้ใช้'))
    doc.add_heading(L('3. Automatic descriptive analysis','3. สถิติเชิงพรรณนาอัตโนมัติ'),1)
    doc.add_paragraph(L('Sample standard deviation uses n−1. Means and quartiles omit missing observations. Outliers use the 1.5×IQR descriptive rule.','ส่วนเบี่ยงเบนมาตรฐานของตัวอย่างใช้ n−1 ค่าเฉลี่ยและควอไทล์ไม่นับข้อมูลขาดหาย ส่วนค่าผิดปกติใช้เกณฑ์พรรณนา 1.5×IQR'))
    if not profile['numeric'].empty:
        basic=['Variable','N','Missing','Mean','Median','Sample SD']
        spread=['Variable','Min','Q1','Q3','Max','IQR outliers']
        doc.add_heading(L('Central tendency and variability','แนวโน้มเข้าสู่ส่วนกลางและความแปรปรวน'),2)
        _add_table(doc,profile['numeric'][basic],max_rows=16)
        doc.add_heading(L('Range and descriptive outliers','ช่วงข้อมูลและค่าผิดปกติเบื้องต้น'),2)
        _add_table(doc,profile['numeric'][spread],max_rows=16)
    else:
        doc.add_paragraph(L('No numerical variables available for summaries.','ไม่มีตัวแปรเชิงตัวเลขสำหรับสรุปผล'))
    fig=_figure_png(df,run)
    if fig:
        doc.add_picture(BytesIO(fig),width=Inches(6.25))
        doc.add_paragraph(L('Illustrative figure based on the current working dataset or final held-out regression predictions.','กราฟตัวอย่างสร้างจากข้อมูลทำงานปัจจุบัน หรือค่าพยากรณ์บนชุดข้อมูลทดสอบที่กันไว้'))
    doc.add_heading(L('4. Correlation screening','4. การสำรวจสหสัมพันธ์'),1)
    doc.add_paragraph(L('Pairwise complete observations. Pearson correlation is descriptive; reported p-values are unadjusted. Multiple comparisons and confounding may be present.','คำนวณจากคู่ข้อมูลที่ไม่ขาดหาย สหสัมพันธ์เพียร์สันเป็นผลพรรณนา ค่า p ยังไม่ปรับการทดสอบหลายครั้ง และอาจมีตัวแปรกวน'))
    _add_table(doc,profile['pairs'],max_rows=12)
    doc.add_heading(L('5. Categorical / survey variables','5. ตัวแปรหมวดหมู่และข้อมูลสำรวจ'),1)
    for col,table in profile['categorical'].items():
        doc.add_heading(str(col),2);_add_table(doc,table,max_rows=10)
    doc.add_heading(L('6. Regression','6. การถดถอย'),1)
    if run is None:doc.add_paragraph(L('No regression analysis was run or retained for the current working dataset.','ยังไม่มีผลการถดถอยสำหรับข้อมูลทำงานชุดปัจจุบัน'))
    else:
        doc.add_paragraph(L(f'Numeric target: {run.target}. Predictors: {", ".join(run.features)}. Chosen model: {run.best_model}. Split: {run.split_method}, test fraction: {run.test_fraction}, random seed: {run.seed}; train n={run.train_n}, held-out test n={run.test_n}.',f'ผลลัพธ์เป้าหมาย: {run.target} · X: {", ".join(run.features)} · แบบจำลอง: {run.best_model} · แบ่งข้อมูล: {run.split_method} · สัดส่วนทดสอบ: {run.test_fraction} · seed: {run.seed} · จำนวนฝึก: {run.train_n} · จำนวนทดสอบ: {run.test_n}'))
        doc.add_paragraph(L('Model selection used cross-validation on the training subset only. The held-out test subset was evaluated once after selection and fitting; imputation and encoding are fit within each training fold.','การเลือกแบบจำลองใช้ cross-validation เฉพาะข้อมูลฝึกเท่านั้น ส่วนข้อมูลทดสอบที่กันไว้ประเมินครั้งเดียวหลังการคัดเลือก โดยขั้นตอนเติมค่าหายและเข้ารหัสปรับจากข้อมูลฝึกในแต่ละรอบ'))
        doc.add_heading(L('Training cross-validation (selection)','การตรวจสอบไขว้ของชุดฝึก (เพื่อเลือกแบบจำลอง)'),2);_add_table(doc,run.cv_metrics)
        doc.add_heading(L('Final held-out test','ผลทดสอบสุดท้ายบนข้อมูลที่กันไว้'),2);_add_table(doc,run.final_metrics)
        for note in run.notes:doc.add_paragraph(note,style='List Bullet')
    doc.add_heading(L('7. Interpretation and limitations','7. การตีความและข้อจำกัด'),1)
    for finding in profile['findings']:
        if lang == 'th':
            if finding.startswith('Computed descriptive summaries'):
                finding=f'คำนวณสถิติให้ตัวแปรเชิงตัวเลข {len(profile["numeric"])} ตัวแปร'
            elif finding.startswith('Largest observed |Pearson r|') and not profile['pairs'].empty:
                highest=profile['pairs'].iloc[0]
                finding=f'คู่ตัวแปรสัมพันธ์สูงสุดในการสำรวจ: {highest["Variable A"]} และ {highest["Variable B"]} (r={highest["r"]:.2f}) ไม่ใช่หลักฐานเหตุและผล'
            elif 'pairs were examined' in finding:
                finding=f'ตรวจสอบ {len(profile["pairs"])} คู่ จึงต้องระวังค่า p จากการทดสอบหลายครั้ง'
            else:
                finding='ตรวจสอบคุณภาพและความหมายของตัวแปรเพิ่มเติม: '+finding
        doc.add_paragraph(finding,style='List Bullet')
    doc.add_paragraph(L('Results are descriptive and predictive, not causal. Verify units, sampling, measurement validity, independence, and data leakage before using results as scientific evidence. Automated interpretations are rule-based and must be reviewed by the researcher.','ผลลัพธ์เป็นการพรรณนาและพยากรณ์ ไม่ยืนยันเหตุและผล ต้องตรวจหน่วย การสุ่มตัวอย่าง ความเที่ยงตรง การพึ่งพากันของข้อมูล และการรั่วไหลของข้อมูลก่อนนำไปใช้อ้างอิงทางวิชาการ คำอธิบายอัตโนมัติมาจากกฎและต้องให้ผู้วิจัยทบทวน'))
    doc.add_paragraph(L('Methods: pandas 2.x, scipy, scikit-learn and Python Plotly/Matplotlib. Numeric/categorical inference and default chart selection are deterministic. Full export tables and cleaned CSV are available separately.','เครื่องมือ: pandas, scipy, scikit-learn และ Plotly/Matplotlib การอนุมานประเภทตัวแปรและการแนะนำกราฟใช้กฎแน่นอน สามารถดาวน์โหลดตารางฉบับเต็มและ CSV แยกต่างหาก'))
    out=BytesIO();doc.save(out);return out.getvalue()
