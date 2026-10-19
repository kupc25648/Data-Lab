from __future__ import annotations

from io import BytesIO
from typing import Any
from zipfile import ZipFile, ZIP_DEFLATED

import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn


def csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")


def tables_zip(tables: dict[str, pd.DataFrame]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as zf:
        for name, df in tables.items():
            safe = name.lower().replace(" ", "_").replace("/", "_")
            zf.writestr(f"{safe}.csv", csv_bytes(df))
    buffer.seek(0)
    return buffer.read()


def _set_defaults(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)


def _add_df(doc: Document, df: pd.DataFrame, max_rows: int = 20) -> None:
    if df is None or df.empty:
        doc.add_paragraph("No data available.")
        return
    show = df.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(show.columns))
    table.style = "Table Grid"
    for i, col in enumerate(show.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in show.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(show.columns):
            val = row[col]
            if isinstance(val, float):
                cells[i].text = f"{val:.4g}"
            else:
                cells[i].text = str(val)


def research_summary_docx(title: str, sections: dict[str, str], tables: dict[str, pd.DataFrame]) -> bytes:
    doc = Document()
    _set_defaults(doc)
    doc.add_heading(title, level=0)
    for sec, text in sections.items():
        doc.add_heading(sec, level=1)
        for para in str(text).split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
    if tables:
        doc.add_heading("Appendix Tables", level=1)
        for name, df in tables.items():
            doc.add_heading(name, level=2)
            _add_df(doc, df)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
